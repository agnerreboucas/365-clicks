"""Gera pages/calendario-365.js a partir da planilha do calendário editorial.

Uso: python3 tools/importar_calendario.py [conteudo/365_Clicks_Calendario_Editorial_365_Dias.xlsx]
Requer: pip install openpyxl
"""
import json
import sys

import openpyxl

ORIGEM = sys.argv[1] if len(sys.argv) > 1 else "conteudo/365_Clicks_Calendario_Editorial_365_Dias.xlsx"
DESTINO = "pages/calendario-365.js"

wb = openpyxl.load_workbook(ORIGEM, data_only=True)
linhas = [list(r) for r in wb["365 Desafios"].iter_rows(min_row=2, values_only=True) if r[0]]

# ---------------------------------------------------------------------------
# Regras de compatibilidade da Matriz Criativa (definidas pela curadoria).
# Colunas: 7 técnica, 8 enquadramento, 9 linguagem, 10 equipamento.
# A planilha original não é alterada; as mudanças saem em conteudo/ajustes_calendario.md.
# ---------------------------------------------------------------------------
REGRAS = [
    ("Close vira “Close ou plano detalhe”: close para pessoas, plano detalhe para objetos.",
     lambda r: r[8] == "Close", lambda r: r.__setitem__(8, "Close ou plano detalhe")),
    ("Fotojornalismo não usa flash: o equipamento Flash vira o recurso “Foco e desfoque”.",
     lambda r: r[9] == "Fotojornalística" and r[10] == "Flash", lambda r: r.__setitem__(10, "Foco e desfoque")),
]
ajustes = []
for descricao, condicao, aplicar in REGRAS:
    dias_afetados = []
    for r in linhas:
        if condicao(r):
            antes = " + ".join(str(x) for x in r[7:11])
            aplicar(r)
            dias_afetados.append((r[0], r[1], r[3], antes, " + ".join(str(x) for x in r[7:11])))
    ajustes.append((descricao, dias_afetados))

# Pontos para a curadoria revisar (não alterados automaticamente)
revisar = [(r[0], r[1], r[3]) for r in linhas if r[9] == "Documental" and r[7] == "Dupla exposição"]
listas = {k: [] for k in ("ciclos", "objetivos", "tecnicas", "enquadramentos", "linguagens", "equipamentos")}


def indice(nome, valor):
    lista = listas[nome]
    if valor not in lista:
        lista.append(valor)
    return lista.index(valor)


dias = [[r[3], indice("ciclos", r[2]), indice("objetivos", r[4]), indice("tecnicas", r[7]),
         indice("enquadramentos", r[8]), indice("linguagens", r[9]), indice("equipamentos", r[10])] for r in linhas if r[0]]
dados = dict(listas, editorias=[list(r) for r in wb["Linha Editorial"].iter_rows(min_row=2, values_only=True) if r[0]], dias=dias)

with open(DESTINO, "w", encoding="utf-8") as f:
    f.write("/* 365 Clicks — calendário editorial dos 365 desafios\n"
            "   Gerado por tools/importar_calendario.py a partir de " + ORIGEM + ".\n"
            "   Cada dia: [tema, ciclo, objetivo, técnica, enquadramento, linguagem, equipamento] (índices nas listas).\n"
            "   Não edite à mão: altere a planilha e rode o script de novo. */\n")
    f.write("window.CALENDARIO_365 = " + json.dumps(dados, ensure_ascii=False, separators=(",", ":")) + ";\n")
with open("conteudo/ajustes_calendario.md", "w", encoding="utf-8") as f:
    f.write("# Ajustes aplicados ao calendário\n\nGerado por `tools/importar_calendario.py`. A planilha original não foi alterada.\n")
    for descricao, dias_afetados in ajustes:
        f.write(f"\n## {descricao}\n\n{len(dias_afetados)} dias.\n\n| Dia | Data | Tema | Antes | Depois |\n|---|---|---|---|---|\n")
        for d in dias_afetados:
            f.write(f"| {d[0]} | {d[1]} | {d[2]} | {d[3]} | {d[4]} |\n")
    f.write("\n## Para a curadoria revisar\n\nLinguagem **Documental** com técnica **Dupla exposição** (a dupla exposição altera a cena; o documental pede registro fiel). "
            f"{len(revisar)} dias, sem mudança automática:\n\n")
    for d in revisar:
        f.write(f"- Dia {d[0]} ({d[1]}): {d[2]}\n")
print(f"{len(dias)} dias gravados em {DESTINO}; ajustes em conteudo/ajustes_calendario.md")
